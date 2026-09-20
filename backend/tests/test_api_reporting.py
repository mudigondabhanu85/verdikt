import csv
import io
import uuid

from app.models.finding import Finding
from app.models.scan import AgentJob, ScanRun
from tests.conftest import create_project_and_version, register_org_admin, session_scope


async def _seed_scan_run_with_findings(
    db_adapter, version_id: str, *, findings: list[dict], status: str = "completed"
) -> str:
    async with session_scope(db_adapter) as session:
        scan_run = ScanRun(version_id=uuid.UUID(version_id), status=status, requested_by=uuid.uuid4())
        session.add(scan_run)
        await session.commit()
        await session.refresh(scan_run)

        job = AgentJob(scan_run_id=scan_run.id, agent_type="header_config", status="completed")
        session.add(job)
        await session.commit()
        await session.refresh(job)

        for spec in findings:
            session.add(
                Finding(
                    scan_run_id=scan_run.id,
                    agent_job_id=job.id,
                    check_id=spec["check_id"],
                    title=spec.get("title", "Test Finding"),
                    severity=spec.get("severity", "Medium"),
                    owasp_2025_category="A02 Security Misconfiguration",
                    cwe_id="CWE-693",
                    cvss_vector="AV:N/AC:L/PR:N/UI:N/S:U/C:N/I:N/A:N",
                    cvss_score=5.0,
                    affected_endpoints=spec.get("affected_endpoints", ["http://site.test/"]),
                    plain_language_summary="summary text",
                    technical_description="technical text",
                    remediation="remediation text",
                    retest_status=spec.get("retest_status", "open"),
                    steps_to_reproduce=spec.get("steps_to_reproduce", []),
                )
            )
        await session.commit()
        return str(scan_run.id)


async def test_report_csv_contains_one_row_per_finding(client, db_adapter):
    admin = await register_org_admin(client)
    _, version_id = await create_project_and_version(client, admin["headers"])
    scan_run_id = await _seed_scan_run_with_findings(
        db_adapter,
        version_id,
        findings=[
            {"check_id": "missing-hsts", "severity": "Low"},
            {"check_id": "sql-injection", "severity": "Critical", "affected_endpoints": ["http://site.test/search?q=1"]},
        ],
    )

    resp = await client.get(f"/scan-runs/{scan_run_id}/report.csv", headers=admin["headers"])
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert "attachment" in resp.headers["content-disposition"]

    rows = list(csv.reader(io.StringIO(resp.text)))
    assert rows[0] == [
        "check_id",
        "title",
        "severity",
        "owasp_2025_category",
        "cwe_id",
        "cvss_score",
        "cvss_vector",
        "affected_endpoints",
        "confirmation_status",
        "retest_status",
        "plain_language_summary",
        "remediation",
    ]
    assert len(rows) == 3  # header + 2 findings
    check_ids = {row[0] for row in rows[1:]}
    assert check_ids == {"missing-hsts", "sql-injection"}


async def test_report_vgs_docx_generates_without_a_curated_draft(client, db_adapter):
    """The Reports tab's per-scan-run VGS export must work immediately
    after a scan completes — no VgsReportDraft, no manual curation
    through the VGS workspace, unlike the real workspace's own
    equivalent (GET /versions/{id}/vgs-report-draft/report.docx)."""
    admin = await register_org_admin(client)
    _, version_id = await create_project_and_version(client, admin["headers"])
    scan_run_id = await _seed_scan_run_with_findings(
        db_adapter,
        version_id,
        findings=[
            {"check_id": "missing-hsts", "severity": "Low"},
            {"check_id": "sql-injection", "severity": "Critical", "affected_endpoints": ["http://site.test/search?q=1"]},
        ],
    )

    resp = await client.get(f"/scan-runs/{scan_run_id}/report.vgs.docx", headers=admin["headers"])
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment" in resp.headers["content-disposition"]
    assert len(resp.content) > 0

    # Nothing curated should have been persisted by a one-shot export.
    from sqlalchemy import select as _select

    from app.models.vgs_vulnerability import VgsReportDraft as _VgsReportDraft

    async with session_scope(db_adapter) as session:
        drafts = (
            await session.execute(_select(_VgsReportDraft).where(_VgsReportDraft.version_id == uuid.UUID(version_id)))
        ).scalars().all()
        assert drafts == []


async def test_report_vgs_docx_groups_findings_by_check_and_title(client, db_adapter):
    admin = await register_org_admin(client)
    _, version_id = await create_project_and_version(client, admin["headers"])
    scan_run_id = await _seed_scan_run_with_findings(
        db_adapter,
        version_id,
        findings=[
            {"check_id": "missing-hsts", "severity": "Low", "affected_endpoints": ["http://site.test/a"]},
            {"check_id": "missing-hsts", "severity": "Low", "affected_endpoints": ["http://site.test/b"]},
        ],
    )

    resp = await client.get(f"/scan-runs/{scan_run_id}/report.vgs.docx", headers=admin["headers"])
    assert resp.status_code == 200, resp.text

    from docx import Document

    document = Document(io.BytesIO(resp.content))
    full_text = "\n".join(p.text for p in document.paragraphs)
    # Grouped into one vulnerability *detail section*, not one per
    # endpoint instance — "CVSS Score:" is only emitted once per detail
    # section, so two would mean grouping silently regressed to one row
    # per endpoint. Both endpoints still show up in the merged
    # description text either way.
    assert full_text.count("CVSS Score:") == 1
    assert "site.test/a" in full_text
    assert "site.test/b" in full_text


async def test_report_vgs_docx_includes_steps_to_reproduce_per_endpoint(client, db_adapter):
    """§ items 4/5: the one-shot per-scan-run VGS export must carry each
    finding's real steps_to_reproduce verbatim, with one evidence step
    per affected endpoint (not just the first one found) — a developer
    should be able to tell exactly how to reproduce each confirmed
    instance of the vulnerability."""
    admin = await register_org_admin(client)
    _, version_id = await create_project_and_version(client, admin["headers"])
    scan_run_id = await _seed_scan_run_with_findings(
        db_adapter,
        version_id,
        findings=[
            {
                "check_id": "missing-hsts",
                "severity": "Low",
                "affected_endpoints": ["http://site.test/a"],
                "steps_to_reproduce": ["1. Request http://site.test/a over HTTPS.", "2. Observe no HSTS header."],
            },
            {
                "check_id": "missing-hsts",
                "severity": "Low",
                "affected_endpoints": ["http://site.test/b"],
                "steps_to_reproduce": ["1. Request http://site.test/b over HTTPS.", "2. Observe no HSTS header."],
            },
        ],
    )

    resp = await client.get(f"/scan-runs/{scan_run_id}/report.vgs.docx", headers=admin["headers"])
    assert resp.status_code == 200, resp.text

    from docx import Document

    document = Document(io.BytesIO(resp.content))
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "1. Request http://site.test/a over HTTPS." in full_text
    assert "1. Request http://site.test/b over HTTPS." in full_text
    assert full_text.count("Step 1 of 2:") == 1
    assert full_text.count("Step 2 of 2:") == 1


async def test_diff_report_classifies_new_fixed_and_still_open(client, db_adapter):
    admin = await register_org_admin(client)
    _, version_id = await create_project_and_version(client, admin["headers"])

    earlier_id = await _seed_scan_run_with_findings(
        db_adapter,
        version_id,
        findings=[
            {"check_id": "missing-hsts", "affected_endpoints": ["http://site.test/"]},  # will be fixed
            {"check_id": "missing-csp", "affected_endpoints": ["http://site.test/"]},  # still open
        ],
    )
    later_id = await _seed_scan_run_with_findings(
        db_adapter,
        version_id,
        findings=[
            {"check_id": "missing-csp", "affected_endpoints": ["http://site.test/"]},  # still open
            {"check_id": "missing-x-frame-options", "affected_endpoints": ["http://site.test/"]},  # new
        ],
    )

    resp = await client.get(f"/scan-runs/{later_id}/diff/{earlier_id}", headers=admin["headers"])
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["earlier_scan_run_id"] == earlier_id
    assert body["later_scan_run_id"] == later_id
    assert {f["check_id"] for f in body["new_findings"]} == {"missing-x-frame-options"}
    assert {f["check_id"] for f in body["fixed_findings"]} == {"missing-hsts"}
    assert {f["check_id"] for f in body["still_open_findings"]} == {"missing-csp"}


async def test_diff_report_rejects_scan_runs_from_different_versions(client, db_adapter):
    admin = await register_org_admin(client)
    _, version_id_a = await create_project_and_version(client, admin["headers"], version_name="v1")
    _, version_id_b = await create_project_and_version(client, admin["headers"], version_name="v2")

    scan_run_a = await _seed_scan_run_with_findings(db_adapter, version_id_a, findings=[{"check_id": "missing-hsts"}])
    scan_run_b = await _seed_scan_run_with_findings(db_adapter, version_id_b, findings=[{"check_id": "missing-hsts"}])

    resp = await client.get(f"/scan-runs/{scan_run_b}/diff/{scan_run_a}", headers=admin["headers"])
    assert resp.status_code == 400


async def test_dashboard_aggregates_open_findings_and_recent_scan_runs(client, db_adapter):
    admin = await register_org_admin(client)
    _, version_id = await create_project_and_version(client, admin["headers"])

    await _seed_scan_run_with_findings(
        db_adapter,
        version_id,
        findings=[
            {"check_id": "missing-hsts", "severity": "Low", "retest_status": "open"},
            {"check_id": "sql-injection", "severity": "Critical", "retest_status": "open"},
            {"check_id": "missing-csp", "severity": "Medium", "retest_status": "fixed"},
        ],
    )

    resp = await client.get("/organizations/me/dashboard", headers=admin["headers"])
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["total_projects"] == 1
    assert body["total_versions"] == 1
    assert body["total_scan_runs"] == 1
    # "fixed" finding is excluded from open counts
    assert body["open_findings_by_severity"]["Critical"] == 1
    assert body["open_findings_by_severity"]["Low"] == 1
    assert body["open_findings_by_severity"]["Medium"] == 0
    assert len(body["recent_scan_runs"]) == 1
    assert body["recent_scan_runs"][0]["finding_counts_by_severity"]["Critical"] == 1


async def test_dashboard_only_counts_callers_own_org(client, db_adapter):
    admin_a = await register_org_admin(client, org_name="Org A", email="a@example.io")
    _, version_id_a = await create_project_and_version(client, admin_a["headers"])
    await _seed_scan_run_with_findings(db_adapter, version_id_a, findings=[{"check_id": "missing-hsts"}])

    admin_b = await register_org_admin(client, org_name="Org B", email="b@example.io")
    await create_project_and_version(client, admin_b["headers"])

    resp = await client.get("/organizations/me/dashboard", headers=admin_b["headers"])
    assert resp.status_code == 200
    body = resp.json()
    assert body["total_projects"] == 1  # only Org B's own project, not Org A's
    assert body["total_scan_runs"] == 0
    assert body["recent_scan_runs"] == []
